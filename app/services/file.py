import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from app.models.bill import Bill
from app.models.category import Category
from app.models.invoice import Invoice
from app.models.transaction import Transaction


class PDFReportGenerator:
    def __init__(self):
        pass

    def generate_invoice_pdf(self, invoice: Invoice) -> bytes:
        buffer = io.BytesIO()
        p = canvas.Canvas(buffer, pagesize=A4)

        # Header
        p.setFont("Helvetica-Bold", 16)
        p.drawString(100, 750, f"Invoice #{invoice.invoice_number}")

        p.setFont("Helvetica", 12)
        p.drawString(100, 720, f"Date: {invoice.issue_date}")
        p.drawString(100, 700, f"Due Date: {invoice.due_date}")
        p.drawString(100, 680, f"Customer: {invoice.customer.name}")

        # Items
        y_position = 620
        p.setFont("Helvetica-Bold", 12)
        p.drawString(100, y_position, "Description")
        p.drawString(300, y_position, "Quantity")
        p.drawString(400, y_position, "Price")
        p.drawString(500, y_position, "Total")

        y_position -= 20
        p.setFont("Helvetica", 12)

        for item in invoice.items:
            p.drawString(100, y_position, item.description)
            p.drawString(300, y_position, str(item.quantity))
            p.drawString(400, y_position, f"${item.price}")
            p.drawString(500, y_position, f"${item.total}")
            y_position -= 20

        # Total
        y_position -= 20
        p.setFont("Helvetica-Bold", 14)
        p.drawString(400, y_position, f"Total: ${invoice.total}")

        p.save()
        buffer.seek(0)
        return buffer.getvalue()

    def generate_bill_pdf(self, bill: Bill) -> bytes:
        buffer = io.BytesIO()
        p = canvas.Canvas(buffer, pagesize=A4)

        # Header
        p.setFont("Helvetica-Bold", 16)
        p.drawString(100, 750, f"Bill #{bill.bill_number}")

        p.setFont("Helvetica", 12)
        p.drawString(100, 720, f"Date: {bill.issue_date}")
        p.drawString(100, 700, f"Due Date: {bill.due_date}")
        p.drawString(100, 680, f"Vendor: {bill.vendor.name}")

        # Items
        y_position = 620
        p.setFont("Helvetica-Bold", 12)
        p.drawString(100, y_position, "Description")
        p.drawString(300, y_position, "Quantity")
        p.drawString(400, y_position, "Price")
        p.drawString(500, y_position, "Total")

        y_position -= 20
        p.setFont("Helvetica", 12)

        for item in bill.items:
            p.drawString(100, y_position, item.description)
            p.drawString(300, y_position, str(item.quantity))
            p.drawString(400, y_position, f"${item.price}")
            p.drawString(500, y_position, f"${item.total}")
            y_position -= 20

        # Total
        y_position -= 20
        p.setFont("Helvetica-Bold", 14)
        p.drawString(400, y_position, f"Total: ${bill.total}")

        p.save()
        buffer.seek(0)
        return buffer.getvalue()


class ExcelExporter:
    @staticmethod
    def export_transactions_to_excel(transactions: list[Transaction]) -> bytes:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Transactions"

        # Header
        headers = ["ID", "Date", "Type", "Amount", "Description", "Category"]
        for col, header in enumerate(headers, 1):
            cell = worksheet.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")

        # Data
        for row, transaction in enumerate(transactions, 2):
            worksheet.cell(row=row, column=1, value=transaction.id)
            worksheet.cell(row=row, column=2, value=str(transaction.transferred_at))
            worksheet.cell(row=row, column=3, value=transaction.type)
            worksheet.cell(row=row, column=4, value=float(transaction.amount))
            worksheet.cell(row=row, column=5, value=transaction.description)
            worksheet.cell(row=row, column=6, value=transaction.category.name if transaction.category else "")

        buffer = io.BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    @staticmethod
    def export_categories_to_excel(categories: list[Category]) -> bytes:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Categories"

        # Header
        headers = ["ID", "Name", "Type", "Color", "Icon", "Active"]
        for col, header in enumerate(headers, 1):
            cell = worksheet.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")

        # Data
        for row, category in enumerate(categories, 2):
            worksheet.cell(row=row, column=1, value=category.id)
            worksheet.cell(row=row, column=2, value=category.name)
            worksheet.cell(row=row, column=3, value=category.type)
            worksheet.cell(row=row, column=4, value=category.color)
            worksheet.cell(row=row, column=5, value=category.icon)
            worksheet.cell(row=row, column=6, value=category.is_active)

        buffer = io.BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    @staticmethod
    def import_transactions_from_excel(file_bytes: bytes) -> list[dict]:
        workbook = Workbook(io.BytesIO(file_bytes))
        worksheet = workbook.active

        transactions = []
        for row in worksheet.iter_rows(min_row=2, values_only=True):
            transactions.append({
                "date": row[1],
                "type": row[2],
                "amount": row[3],
                "description": row[4],
                "category_id": row[5] if row[5] else None,
            })

        return transactions
